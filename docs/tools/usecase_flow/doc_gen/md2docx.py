#!/usr/bin/env python3
"""Chuyển file đặc tả SkyGo Use Case (Markdown, theo STANDARD.md) sang DOCX.

Dùng:
  python md2docx.py <file.md | thư mục> [...] [--out dist] [--template uc-template.docx]
  python md2docx.py f3-lam-quiz --check-only
  python md2docx.py BKG-UC-001_x.md --assets docs/build/usecase-flow/booking/BKG-UC-001

Việc script làm:
  1. Đọc file <MODULE>-UC-NNN_*.md (bỏ qua README.md, rules.md).
  2. Kiểm tra: header đủ 6 dòng đúng thứ tự, số bước liền mạch, nhánh N.k rẽ từ bước N có thật,
     mã trong {{ref: ...}} có định nghĩa ở 00-module.md hoặc rules.md (cảnh báo, không chặn).
  3. Xoá mọi khối {{ref: ...}}, gom thành MỘT bảng duy nhất và ghi ra docx theo style của template.
     Template là file docx mẫu của trường (cần style "Heading 3" và "Text_in_table").
"""
import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, Twips

REF_RE = re.compile(r"\s*\{\{ref:[^}]*\}\}")
STEP_RE = re.compile(r"^(\d+(?:\.\d+)*)\.\s")
HEADER_LABELS = ["Use case:", "Mục đích:", "Mô tả:", "Actor:", "Tiền điều kiện:", "Hậu điều kiện:"]
COL_W = (4914, 5245)  # độ rộng cột (twips) như bản mẫu của trường
GRID_W = (4503, 4851)
UC_FILE_RE = re.compile(r"^[A-Z][A-Z0-9]*-UC-\d{3}_.+\.md$")
UC_CODE_RE = re.compile(r"\b[A-Z][A-Z0-9]*-UC-\d{3}\b")
BR_CODE_RE = re.compile(r"\bBR-[A-Z][A-Z0-9]*-\d{3}\b")


# ---------- parse markdown ----------
def split_row(line):
    inner = line.strip()
    if inner.startswith("|"):
        inner = inner[1:]
    if inner.endswith("|"):
        inner = inner[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", inner)]


def is_sep(cells):
    return all(re.fullmatch(r":?-{3,}:?", c) for c in cells)


def parse_md(text):
    """Trả về dict: title, desc_title, header (list[(label, value)]), main, alt, diagrams."""
    sections, cur = [], None
    for line in text.splitlines():
        if line.startswith("### "):
            cur = {"title": line[4:].strip(), "lines": []}
            sections.append(cur)
        elif cur is not None:
            cur["lines"].append(line)
    if len(sections) < 3:
        raise ValueError("thiếu mục ### (cần: tên UC, Mô tả use case, Luồng chính, Luồng phụ, Biểu đồ)")

    def table_rows(lines):
        rows = []
        for ln in lines:
            if ln.strip().startswith("|"):
                cells = split_row(ln)
                if not is_sep(cells):
                    rows.append(cells)
        return rows

    out = {"title": sections[0]["title"], "desc_title": sections[1]["title"], "diagrams": []}
    out["header"] = [
        tuple(r)
        for r in table_rows(sections[1]["lines"])
        if r and r[0].replace("**", "") not in ("Trường", "Field")
    ]
    by_title = {s["title"]: s for s in sections[2:]}
    main = table_rows(by_title["Luồng sự kiện chính"]["lines"])
    alt = table_rows(by_title["Luồng sự kiện phụ"]["lines"])
    actor_label = next(
        (r[1].replace("**", "").strip().rstrip(".") for r in out["header"] if r and r[0].replace("**", "") == "Actor:"),
        "Actor",
    )
    column_header = next(
        (r for r in main if len(r) == 2 and re.sub(r"\*", "", r[1]).strip() == "Hệ thống"),
        None,
    )
    out["actor"] = re.sub(r"\*", "", column_header[0]).strip() if column_header else actor_label
    is_column_header = lambda r: (
        len(r) == 2
        and re.sub(r"\*", "", r[0]).strip() == out["actor"].strip()
        and re.sub(r"\*", "", r[1]).strip() == "Hệ thống"
    )
    out["main"] = [r for r in main if not is_column_header(r)]
    out["alt"] = [r for r in alt if not is_column_header(r)]
    dg = by_title.get("Biểu đồ")
    if dg:
        label = None
        for ln in dg["lines"]:
            s = ln.strip()
            label_match = re.match(r"^(Activity|Sequence):", s, re.I)
            if label_match:
                label = label_match.group(1).capitalize() + ":"
                out["diagrams"].append([label, None])
            m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", s)
            if m and out["diagrams"]:
                out["diagrams"][-1][1] = m.group(1)
    return out


# ---------- kiểm tra ----------
def collect_defined(start):
    root = None
    for p in [start.resolve()] + list(start.resolve().parents):
        if (p / "00-module.md").exists():
            root = p
            break
    if root is None:
        return None, None
    brs, ucs = set(), set()
    files = [root / "00-module.md"] + list(root.rglob("rules.md"))
    for f in files:
        t = f.read_text(encoding="utf-8")
        brs |= set(BR_CODE_RE.findall(t))
        ucs |= set(UC_CODE_RE.findall(t))
    return brs, ucs


def check(parsed, raw_text, path, brs, ucs):
    errors, warns = [], []
    labels = [r[0].replace("**", "") for r in parsed["header"]]
    if labels != HEADER_LABELS:
        errors.append(f"header sai hoặc sai thứ tự: {labels}")
    for r in parsed["header"] + parsed["main"] + parsed["alt"]:
        if len(r) != 2:
            errors.append(f"hàng không đủ 2 cột: {r}")
    # bước chính liền mạch
    nums = []
    for r in parsed["main"]:
        for c in r:
            m = STEP_RE.match(REF_RE.sub("", c))
            if m:
                nums.append(m.group(1))
    if any("." in n for n in nums) or [int(n) for n in nums] != list(range(1, len(nums) + 1)):
        errors.append(f"bước luồng chính không liền mạch 1..n: {nums}")
    n_main = len(nums)
    # nhánh phụ
    seen = set()
    for r in parsed["alt"]:
        for c in r:
            m = STEP_RE.match(REF_RE.sub("", c))
            if not m:
                continue
            label = m.group(1)
            parts = label.split(".")
            if len(parts) < 2:
                errors.append(f"bước phụ '{label}' phải có dạng N.k hoặc N.k.m")
                continue
            if not 1 <= int(parts[0]) <= n_main:
                errors.append(f"nhánh '{label}' rẽ từ bước {parts[0]} không tồn tại ở luồng chính")
            if label in seen:
                errors.append(f"trùng số bước '{label}'")
            seen.add(label)
            if len(parts) == 3 and ".".join(parts[:2]) not in seen:
                errors.append(f"bước '{label}' xuất hiện trước đầu nhánh '{'.'.join(parts[:2])}'")
    # ref
    if brs is not None:
        missing = set()
        for blk in re.findall(r"\{\{ref:([^}]*)\}\}", raw_text):
            for c in (x.strip() for x in blk.split(",")):
                if c.startswith("BR-") and c not in brs:
                    missing.add(c)
                elif UC_CODE_RE.fullmatch(c) and c not in ucs:
                    missing.add(c)
                elif not c.startswith("BR-") and not UC_CODE_RE.fullmatch(c) and c != "Suy ra":
                    missing.add(c)
        if missing:
            warns.append("ref chưa có định nghĩa: " + ", ".join(sorted(missing)))
    if "CHG-" in raw_text:
        warns.append("còn mã lịch sử CHG-")
    return errors, warns


# ---------- ghi docx ----------
def clean(text):
    return re.sub(r"\s{2,}", " ", REF_RE.sub("", text)).strip()


def split_sentences(text):
    pieces = re.split(r"(?<=\.)\s+", text)
    out = []
    for p in pieces:
        if out and p and not p[0].isupper():
            out[-1] += " " + p
        else:
            out.append(p)
    return out


def add_runs(par, text):
    for i, part in enumerate(text.split("**")):
        if part:
            run = par.add_run(part)
            if i % 2 == 1:
                run.bold = True


def new_par(cell, first):
    par = cell.paragraphs[0] if first else cell.add_paragraph()
    par.style = "Text_in_table"
    return par


def set_cell(cell, text, *, bold=False, center=False, split=False):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    text = clean(text)
    if bold and "**" not in text:
        text = f"**{text}**" if text else text
    if split:
        chunks = [
            re.sub(r"[•.]\s*$", "", re.sub(r"^\s*•\s*", "", part)).strip()
            for part in re.split(r"<br\s*/?>", text, flags=re.I)
        ]
    else:
        chunks = [p for p in re.split(r"<br\s*/?>", text, flags=re.I)]
    if not chunks:
        chunks = [""]
    for i, ch in enumerate(chunks):
        par = new_par(cell, i == 0)
        if center:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(par, ch)


def style_table(tbl):
    tblPr = tbl._tbl.tblPr
    for child in list(tblPr):
        tblPr.remove(child)
    xml = (
        f'<w:tblPr {nsdecls("w")}>'
        '<w:tblW w:w="0" w:type="auto"/>'
        '<w:tblCellSpacing w:w="15" w:type="dxa"/>'
        '<w:tblBorders>'
        + "".join(
            f'<w:{s} w:val="single" w:sz="2" w:space="0" w:color="auto"/>'
            for s in ("top", "left", "bottom", "right", "insideH", "insideV")
        )
        + "</w:tblBorders>"
        '<w:tblCellMar><w:top w:w="15" w:type="dxa"/><w:left w:w="15" w:type="dxa"/>'
        '<w:bottom w:w="15" w:type="dxa"/><w:right w:w="15" w:type="dxa"/></w:tblCellMar>'
        '<w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" w:firstColumn="1" w:lastColumn="0" w:noHBand="0" w:noVBand="1"/>'
        "</w:tblPr>"
    )
    new = parse_xml(xml)
    tbl._tbl.insert(0, new)
    grid = tbl._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), GRID_W):
        gc.set(qn("w:w"), str(w))


def add_row(tbl, header=False):
    row = tbl.add_row()
    trPr = row._tr.get_or_add_trPr()
    if header:
        trPr.append(OxmlElement("w:tblHeader"))
    sp = OxmlElement("w:tblCellSpacing")
    sp.set(qn("w:w"), "15")
    sp.set(qn("w:type"), "dxa")
    trPr.append(sp)
    for cell, w in zip(row.cells, COL_W):
        cell.width = Twips(w)
    return row


def span_row(tbl, text):
    row = add_row(tbl)
    merged = row.cells[0].merge(row.cells[1])
    set_cell(merged, text, bold=True)


def load_template(path):
    doc = Document(str(path))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith("/image"):
            doc.part.drop_rel(rid)
    return doc


def heading(doc, text):
    p = doc.add_paragraph(text, style="Heading 3")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    return p


def build_docx(parsed, template, base_dir, assets=None, uc_code=None):
    doc = load_template(template)
    heading(doc, parsed["title"])
    heading(doc, parsed["desc_title"])
    tbl = doc.add_table(rows=0, cols=2)
    style_table(tbl)
    for i, (label, value) in enumerate(parsed["header"]):
        row = add_row(tbl, header=(i == 0))
        name = label.replace("**", "")
        set_cell(row.cells[0], name, bold=True)
        set_cell(row.cells[1], value, bold=(i == 0), split=name in ("Tiền điều kiện:", "Hậu điều kiện:"))
    span_row(tbl, "Luồng sự kiện chính")
    row = add_row(tbl)
    set_cell(row.cells[0], parsed["actor"], bold=True, center=True)
    set_cell(row.cells[1], "Hệ thống", bold=True, center=True)
    for left, right in parsed["main"]:
        row = add_row(tbl)
        set_cell(row.cells[0], left)
        set_cell(row.cells[1], right)
    span_row(tbl, "Luồng sự kiện phụ")
    for left, right in parsed["alt"]:
        row = add_row(tbl)
        set_cell(row.cells[0], left)
        set_cell(row.cells[1], right)
    doc.add_paragraph()
    heading(doc, "Biểu đồ")
    for label, img in parsed["diagrams"]:
        p = doc.add_paragraph(label)
        p.paragraph_format.keep_with_next = True
        if not img and assets and uc_code:
            kind = "activity" if label.lower().startswith("activity") else "sequence"
            cand = Path(assets) / f"{uc_code}_{kind}.png"
            if cand.exists():
                doc.add_paragraph().add_run().add_picture(str(cand), width=Inches(6.0))
        if img:
            img_path = (base_dir / img).resolve()
            if img_path.exists():
                doc.add_paragraph().add_run().add_picture(str(img_path), width=Inches(6.0))
            else:
                print(f"  ! không thấy ảnh {img}", file=sys.stderr)
    doc.core_properties.title = parsed["title"]
    return doc


# ---------- main ----------
def find_uc_files(paths):
    files = []
    for p in map(Path, paths):
        if p.is_dir():
            files += sorted(f for f in p.rglob("*.md") if UC_FILE_RE.match(f.name))
        elif UC_FILE_RE.match(p.name):
            files.append(p)
        else:
            print(f"bỏ qua {p} (không phải file <MODULE>-UC-NNN_*.md)", file=sys.stderr)
    return files


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("inputs", nargs="+")
    ap.add_argument("--out", default="dist", help="thư mục xuất docx (mặc định: dist)")
    ap.add_argument("--template", default=str(Path(__file__).with_name("uc-template.docx")))
    ap.add_argument("--assets", help="thư mục chứa <MODULE>-UC-NNN_activity.png / _sequence.png để tự chèn")
    ap.add_argument("--check-only", action="store_true", help="chỉ kiểm tra, không ghi docx")
    args = ap.parse_args()

    files = find_uc_files(args.inputs)
    if not files:
        sys.exit("không có file <MODULE>-UC-NNN_*.md nào")
    failed = 0
    for f in files:
        raw = f.read_text(encoding="utf-8")
        try:
            parsed = parse_md(raw)
        except (ValueError, KeyError) as e:
            print(f"[LỖI] {f.name}: {e}")
            failed += 1
            continue
        brs, ucs = collect_defined(f.parent)
        errors, warns = check(parsed, raw, f, brs, ucs)
        for e in errors:
            print(f"[LỖI] {f.name}: {e}")
        for w in warns:
            print(f"[CẢNH BÁO] {f.name}: {w}")
        if errors:
            failed += 1
            continue
        if args.check_only:
            print(f"[OK] {f.name}")
            continue
        out_dir = Path(args.out)
        out_dir.mkdir(parents=True, exist_ok=True)
        out = out_dir / (f.stem + ".docx")
        m = UC_FILE_RE.match(f.name)
        uc_code = f.stem.split("_", 1)[0] if m else None
        build_docx(parsed, args.template, f.parent, args.assets, uc_code).save(str(out))
        print(f"[OK] {f.name} -> {out}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
